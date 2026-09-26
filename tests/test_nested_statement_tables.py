import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from opendartmcp.excel import dartdoc
from opendartmcp.excel.build_financial_excel import build_workbook
from opendartmcp.excel.verify_workbook import verify
from tests.fixtures import nested_statement_tables_xml


class NestedStatementTablesTest(unittest.TestCase):
    def test_nested_cashflow_table_is_extracted_once(self):
        content = nested_statement_tables_xml()
        model = dartdoc.extract_model(content, dartdoc.CONSOLIDATED)

        self.assertEqual(
            [statement["sheet_name"] for statement in model["statements"]],
            ["연결재무상태표", "연결포괄손익계산서", "연결자본변동표", "연결현금흐름표"],
        )
        cashflow = model["statements"][-1]
        self.assertEqual(len(cashflow["tables"]), 1)
        self.assertNotIn("1,000", str(cashflow["preamble"]))

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested.xlsx"
            build_workbook(model, str(path))
            report = verify(model, str(path), content)
            self.assertTrue(report["ok"], report["failures"])
            workbook = load_workbook(path)
            values = [
                cell.value
                for row in workbook["연결현금흐름표"].iter_rows()
                for cell in row
            ]
            self.assertEqual(values.count(1000), 1)

    def test_nested_wrappers_have_no_depth_limit(self):
        content = nested_statement_tables_xml(depth=2)
        model = dartdoc.extract_model(content, dartdoc.CONSOLIDATED)

        self.assertEqual(model["statements"][-1]["sheet_name"], "연결현금흐름표")
        self.assertEqual(len(model["statements"][-1]["tables"]), 1)

    def test_wrapper_postscript_belongs_to_nested_statement(self):
        content = nested_statement_tables_xml(with_postscripts=True)
        model = dartdoc.extract_model(content, dartdoc.CONSOLIDATED)

        self.assertEqual(
            [statement["postscript"] for statement in model["statements"]],
            [["별첨 주석 참조"]] * 4,
        )


if __name__ == "__main__":
    unittest.main()
